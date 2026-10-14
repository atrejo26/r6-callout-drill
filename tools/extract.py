"""Pull floor plans and callout labels out of the call-sheet deck.

Usage: python3 -I tools/extract.py <deck.pptx> src/

Writes src/maps/<map>_<floor>.webp and src/callouts.json. Each label is stored
as a position normalized to its (cropped, trimmed) floor image, so the app can
place markers at any display size.
"""
import io, json, os, re, sys
from pptx import Presentation
from PIL import Image, ImageOps

A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'

# slide number -> (map id, display name, floor name for each picture in z-order)
MAPS = {
    3:  ('border',     'Border',          ['1F', '2F']),
    4:  ('chalet',     'Chalet',          ['Basement', '1F', '2F']),
    5:  ('clubhouse',  'Clubhouse',       ['Basement', '1F', '2F']),
    6:  ('villa',      'Villa',           ['1F', 'Basement', '2F']),
    7:  ('nighthaven', 'Nighthaven Labs', ['Basement', '1F', '2F']),
    8:  ('fortress',   'Fortress',        ['2F', '1F']),
    9:  ('bank',       'Bank',            ['Basement', '1F', '2F']),
    10: ('kafe',       'Kafe Dostoyevsky', ['1F', '2F', '3F']),
    11: ('lair',       'Lair',            ['2F', '1F', 'Basement']),
}
FLOOR_ORDER = ['Basement', '1F', '2F', '3F']

# Display fixes for typos / formatting in the deck; aliases are extra accepted answers.
FIXES = {
    'Dinning': ('Dining', []),
    'BottomSqr': ('Bottom Sqr', []),
    'DMZ(DeadManZone)': ('DMZ', ['Dead Man Zone']),
    'BLUE HALL': ('Blue Hall', []),
    'Tellers  hall': ('Tellers Hall', []),
    'Detention or Jail': ('Detention', ['Jail']),
    'boost': ('Boost', []),
}


def walk(shapes, xf):
    """Yield (shape, slide-space rect), resolving nested group transforms."""
    for sh in shapes:
        if sh.shape_type == 6:
            x = sh._element.grpSpPr.find(A + 'xfrm')
            off, ext, coff, cext = (x.find(A + n) for n in ('off', 'ext', 'chOff', 'chExt'))
            ox, oy = int(off.get('x')), int(off.get('y'))
            gsx = int(ext.get('cx')) / (int(cext.get('cx')) or 1)
            gsy = int(ext.get('cy')) / (int(cext.get('cy')) or 1)
            cx, cy = int(coff.get('x')), int(coff.get('y'))
            sx, sy, tx, ty = xf
            yield from walk(sh.shapes, (sx * gsx, sy * gsy, tx + sx * (ox - cx * gsx), ty + sy * (oy - cy * gsy)))
        elif sh.left is not None:
            sx, sy, tx, ty = xf
            yield sh, (tx + sx * sh.left, ty + sy * sh.top, sx * sh.width, sy * sh.height)


def label_info(sh):
    paras, size, bold, color = [], 0, False, None
    for para in sh.text_frame.paragraphs:
        t = ''.join(r.text for r in para.runs).strip()
        if t:
            paras.append(t)
        for r in para.runs:
            if r.font.size:
                size = max(size, r.font.size.pt)
            bold = bold or bool(r.font.bold)
            try:
                if r.font.color and r.font.color.type is not None:
                    color = str(r.font.color.rgb)
            except (AttributeError, KeyError):
                pass
    return ' '.join(paras), size, bold, color


def trim_box(img):
    """Bounding box of non-white content, padded, so empty margins don't shrink the plan."""
    g = ImageOps.invert(img.convert('L')).point(lambda v: 255 if v > 20 else 0)
    box = g.getbbox() or (0, 0, *img.size)
    pad = round(0.02 * max(img.size))
    return (max(0, box[0] - pad), max(0, box[1] - pad), min(img.width, box[2] + pad), min(img.height, box[3] + pad))


def main(deck, outdir):
    os.makedirs(os.path.join(outdir, 'maps'), exist_ok=True)
    prs = Presentation(deck)
    maps = []
    for num, (mid, name, floor_names) in MAPS.items():
        pics, texts = [], []
        for z, (sh, rect) in enumerate(walk(prs.slides[num - 1].shapes, (1, 1, 0, 0))):
            if sh.shape_type == 13:
                pics.append((z, sh, rect))
            elif sh.has_text_frame and sh.text_frame.text.strip():
                text, size, bold, color = label_info(sh)
                # Skip slide titles and floor headings (black text, not red callouts).
                if color != 'FF0000':
                    continue
                texts.append((text, rect, sh.rotation, bold))
        assert len(pics) == len(floor_names), (mid, len(pics))
        floors = []
        for (z, sh, rect), fname in zip(pics, floor_names):
            img = Image.open(io.BytesIO(sh.image.blob)).convert('RGB')
            W, H = img.size
            img = img.crop((round(sh.crop_left * W), round(sh.crop_top * H),
                            round(W - sh.crop_right * W), round(H - sh.crop_bottom * H)))
            floors.append(dict(name=fname, z=z, rect=rect, img=img, labels=[]))
        for text, (x, y, w, h), rot, bold in texts:
            cx, cy = x + w / 2, y + h / 2
            hits = [f for f in floors if f['rect'][0] <= cx <= f['rect'][0] + f['rect'][2]
                    and f['rect'][1] <= cy <= f['rect'][1] + f['rect'][3]]
            if not hits:
                continue
            f = max(hits, key=lambda f: f['z'])  # topmost picture wins where plans overlap
            fx, fy, fw, fh = f['rect']
            text = re.sub(r'\s+', ' ', text).strip()
            disp, aliases = FIXES.get(text, (text, []))
            disp = disp[:1].upper() + disp[1:]
            f['labels'].append(dict(name=disp, aliases=aliases, room=bold, rot=round(rot),
                                    u=(cx - fx) / fw, v=(cy - fy) / fh))
        out_floors = []
        for f in sorted(floors, key=lambda f: FLOOR_ORDER.index(f['name'])):
            img = f['img']
            l, t, r, b = trim_box(img)
            W, H = img.size
            img = img.crop((l, t, r, b))
            labels = []
            for lab in f['labels']:
                u = (lab.pop('u') * W - l) / (r - l)
                v = (lab.pop('v') * H - t) / (b - t)
                labels.append(dict(lab, x=round(u, 4), y=round(v, 4)))
            fn = f"{mid}_{f['name'].lower()}.webp"
            img.save(os.path.join(outdir, 'maps', fn), 'WEBP', quality=82, method=6)
            out_floors.append(dict(name=f['name'], image=f'maps/{fn}', w=img.width, h=img.height, callouts=labels))
            print(f"{mid:11s} {f['name']:8s} {img.width}x{img.height} {len(labels)} callouts")
        maps.append(dict(id=mid, name=name, floors=out_floors))
    with open(os.path.join(outdir, 'callouts.json'), 'w') as fh:
        json.dump(dict(source='Paltry\'s R6 Map Call Sheets', maps=maps), fh, indent=1)


if __name__ == '__main__':
    main(*sys.argv[1:3])
